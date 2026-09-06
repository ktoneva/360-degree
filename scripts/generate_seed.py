"""Generates supabase/migrations/20260905233003_seed_item_bank.sql from the
spec spreadsheet. Run once when the item bank changes; the output is a plain
SQL file checked into migrations, not run at build time.
"""
import io
import re
import openpyxl

SRC = r"C:\Users\ktone\Downloads\360_educational_leaders_v7.xlsx"
OUT = r"C:\Users\ktone\Documents\360-educational-leaders\supabase\migrations\20260905233003_seed_item_bank.sql"

RATER_COLUMNS = [
    ("Self", "self"),
    ("Manager", "manager"),
    ("Peers", "peer"),
    ("Direct reports", "direct_report"),
    ("Others", "other"),
]


def sql_str(value):
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def rater_array(flags):
    groups = [g for present, g in flags if present]
    if not groups:
        raise ValueError("item has no rater groups")
    inner = ", ".join(sql_str(g) for g in groups)
    return f"ARRAY[{inner}]::rater_group[]"


def competency_subquery(number, variant):
    return f"(select id from competencies where number = {number} and variant = {sql_str(variant)})"


def equivalent_item_subquery(item_number):
    return (
        "(select i.id from items i join competencies c on c.id = i.competency_id "
        f"where c.number = 9 and c.variant = 'standard' and i.item_number = {item_number})"
    )


def item_values_line(comp_subquery, item_number, behaviour, groups_sql, response_type, is_integrity, design_note, equivalent_sql="null"):
    return "  ({comp}, {num}, {text}, {groups}, {rtype}, {integrity}, {equiv}, {note})".format(
        comp=comp_subquery,
        num=item_number,
        text=sql_str(behaviour),
        groups=groups_sql,
        rtype=sql_str(response_type),
        integrity=str(is_integrity).lower(),
        equiv=equivalent_sql,
        note=sql_str(design_note),
    )


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    ws = wb["Item bank"]
    rows = list(ws.iter_rows(values_only=True))
    header, rows = rows[0], rows[1:]
    col = {name: i for i, name in enumerate(header)}

    competencies = {}  # number -> name
    standard_items = []

    for row in rows:
        if row[col["Competency"]] is None:
            continue
        comp_label = row[col["Competency"]]
        m = re.match(r"(\d+)\.\s*(.+)", comp_label)
        comp_number, comp_name = int(m.group(1)), m.group(2).strip()
        competencies[comp_number] = comp_name

        item_number = row[col["No."]]
        behaviour = row[col["Behaviour"]]
        design_note = row[col["Design note"]]
        is_integrity = bool(design_note and "INTEGRITY ITEM" in design_note)
        flags = [(row[col[sheet_col]] == "\u25cf", group) for sheet_col, group in RATER_COLUMNS]

        standard_items.append(dict(
            comp_number=comp_number,
            item_number=item_number,
            behaviour=behaviour,
            flags=flags,
            design_note=design_note,
            is_integrity=is_integrity,
        ))

    comp9_groups_by_number = {
        it["item_number"]: it["flags"] for it in standard_items if it["comp_number"] == 9
    }

    ops_ws = wb["9-OPS variant"]
    ops_rows = list(ops_ws.iter_rows(values_only=True))
    ops_data_rows = [r for r in ops_rows[4:] if r[0] is not None]

    out = io.StringIO()
    out.write("-- Seed data for the item bank: 9 standard competencies (90 items) plus\n")
    out.write("-- the 9-OPS variant (10 items), generated from 360_educational_leaders_v7.xlsx.\n")
    out.write("-- Regenerate with scripts/generate_seed.py if the spec changes.\n\n")

    out.write("insert into competencies (number, variant, name) values\n")
    comp_lines = [f"  ({n}, 'standard', {sql_str(name)})" for n, name in sorted(competencies.items())]
    comp_lines.append(f"  (9, 'ops', {sql_str('Operational standards and service quality')})")
    out.write(",\n".join(comp_lines))
    out.write(";\n\n")

    out.write("insert into items (competency_id, item_number, behaviour_text, asked_rater_groups, response_type, is_integrity_item, equivalent_item_id, design_note) values\n")
    lines = []
    for it in standard_items:
        response_type = "yes_no_not_observed" if it["is_integrity"] else "scale"
        lines.append(item_values_line(
            comp_subquery=competency_subquery(it["comp_number"], "standard"),
            item_number=it["item_number"],
            behaviour=it["behaviour"],
            groups_sql=rater_array(it["flags"]),
            response_type=response_type,
            is_integrity=it["is_integrity"],
            design_note=it["design_note"],
        ))
    out.write(",\n".join(lines))
    out.write(";\n\n")

    out.write("-- 9-OPS variant: same architecture, non-teaching wording. Rater-group\n")
    out.write("-- visibility is inherited from the mapped standard competency-9 item, since\n")
    out.write("-- the source sheet doesn't give ops items their own per-item breakdown.\n")
    out.write("insert into items (competency_id, item_number, behaviour_text, asked_rater_groups, response_type, is_integrity_item, equivalent_item_id, design_note) values\n")
    ops_lines = []
    for row in ops_data_rows:
        item_number, behaviour, maps_to = row[0], row[1], row[2]
        groups = comp9_groups_by_number[item_number]
        ops_lines.append(item_values_line(
            comp_subquery=competency_subquery(9, "ops"),
            item_number=item_number,
            behaviour=behaviour,
            groups_sql=rater_array(groups),
            response_type="scale",
            is_integrity=False,
            design_note=f"Maps to teaching version {maps_to}.",
            equivalent_sql=equivalent_item_subquery(item_number),
        ))
    out.write(",\n".join(ops_lines))
    out.write(";\n")

    wb.close()
    with io.open(OUT, "w", encoding="utf-8") as f:
        f.write(out.getvalue())

    print("wrote", OUT)
    print(len(standard_items), "standard items,", len(ops_data_rows), "ops items")


if __name__ == "__main__":
    main()
